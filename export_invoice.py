"""Excel and PDF invoice export."""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.page import PageMargins
from reportlab.lib.colors import HexColor, white
from reportlab.lib.pagesizes import A4
from reportlab.lib.utils import ImageReader
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen import canvas

from invoice_model import (
    ISSUER,
    LOGO_PATH,
    SELLER,
    Invoice,
    amount_in_words,
    fmt_money,
    fmt_qty,
    footer_lines,
    meta_lines,
)

NAVY = "0E2A47"
INK = "1C2834"
MUTED = "5C6B7A"
LINE = "D9E0E8"
ZEBRA = "F5F7FB"
TOTAL = "E6E8FF"
WHITE = "FFFFFF"

PAGE_W, PAGE_H = A4
MARGIN = 42


def _font(name: str, size: int, bold: bool = False, italic: bool = False, color: str = INK) -> Font:
    return Font(name=name, size=size, bold=bold, italic=italic, color=color)


def _fill(color: str) -> PatternFill:
    return PatternFill("solid", fgColor=color)


def _side(color: str = LINE, style: str = "thin") -> Side:
    return Side(style=style, color=color)


def _border(bottom: bool = False, top: bool = False) -> Border:
    return Border(
        left=_side("FFFFFF", "thin"),
        right=_side("FFFFFF", "thin"),
        top=_side(NAVY if top else "FFFFFF"),
        bottom=_side(LINE if bottom else "FFFFFF"),
    )


def _apply(cell, font=None, align=None, fill=None, border=None, fmt=None):
    if font is not None:
        cell.font = font
    if align is not None:
        cell.alignment = align
    if fill is not None:
        cell.fill = fill
    if border is not None:
        cell.border = border
    if fmt is not None:
        cell.number_format = fmt
    return cell


def write_excel(invoice: Invoice, path: Path) -> None:
    profile = invoice.profile
    wb = Workbook()
    ws = wb.active
    ws.title = "Račun"

    ws.page_setup.orientation = "portrait"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.page_setup.horizontalCentered = True
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins = PageMargins(left=0.45, right=0.45, top=0.4, bottom=0.4, header=0.15, footer=0.15)
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 120
    ws.sheet_properties.tabColor = NAVY
    ws.oddFooter.center.text = f"Veritas Software  ·  Račun br. {invoice.number}"
    ws.oddFooter.center.font = "Calibri"
    ws.oddFooter.center.size = 9

    widths = {"A": 3, "B": 52, "C": 14, "D": 18, "E": 20, "F": 3}
    for column, width in widths.items():
        ws.column_dimensions[column].width = width

    for col in range(1, 6):
        _apply(ws.cell(1, col, ""), fill=_fill(NAVY))
    ws.row_dimensions[1].height = 8

    ws.row_dimensions[2].height = 36
    ws.row_dimensions[3].height = 8
    _apply(
        ws.cell(2, 5, f"Račun br. {invoice.number}"),
        font=_font("Calibri", 22, bold=True, color=NAVY),
        align=Alignment(horizontal="right", vertical="center"),
    )
    if LOGO_PATH.exists():
        image = XLImage(str(LOGO_PATH))
        image.width = 188
        image.height = 33
        ws.add_image(image, "B2")

    label_font = _font("Calibri", 8, bold=True, color="7A8794")
    name_font = _font("Calibri", 12, bold=True, color=NAVY)
    body_font = _font("Calibri", 10, color=INK)
    muted_font = _font("Calibri", 10, color=MUTED)
    bold_font = _font("Calibri", 10, bold=True, color=INK)
    left = Alignment(horizontal="left", vertical="center", wrap_text=True)
    right = Alignment(horizontal="right", vertical="center", wrap_text=True)

    ws.row_dimensions[5].height = 14
    _apply(ws.cell(5, 2, "IZDAVALAC"), font=label_font, align=left)
    _apply(ws.cell(5, 4, "KUPAC"), font=label_font, align=left)

    ws.merge_cells("B6:C6")
    ws.merge_cells("D6:E6")
    _apply(ws.cell(6, 2, SELLER.name), font=name_font, align=left)
    _apply(ws.cell(6, 4, invoice.buyer.name), font=name_font, align=left)
    ws.row_dimensions[6].height = 34

    rows_info = (
        (7, SELLER.address, invoice.buyer.address, body_font),
        (8, f"Matični broj: {SELLER.registration}", f"Matični broj: {invoice.buyer.registration}", muted_font),
        (9, f"Poreski broj: {SELLER.tax}", f"Poreski broj: {invoice.buyer.tax}", muted_font),
    )
    for row, left_text, right_text, font in rows_info:
        ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=3)
        ws.merge_cells(start_row=row, start_column=4, end_row=row, end_column=5)
        _apply(ws.cell(row, 2, left_text), font=font, align=left)
        _apply(ws.cell(row, 4, right_text), font=font, align=left)
        ws.row_dimensions[row].height = 18

    ws.row_dimensions[10].height = 10
    for offset, line in enumerate(profile.bank_lines):
        row = 11 + offset
        ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=3)
        _apply(ws.cell(row, 2, line), font=muted_font, align=left)
        ws.row_dimensions[row].height = 32 if offset == len(profile.bank_lines) - 1 else 18

    for offset, (text, is_bold) in enumerate(meta_lines(invoice)):
        row = 11 + offset
        ws.merge_cells(start_row=row, start_column=4, end_row=row, end_column=5)
        _apply(ws.cell(row, 4, text), font=bold_font if is_bold else muted_font, align=left)
        ws.row_dimensions[row].height = max(ws.row_dimensions[row].height or 0, 18)

    header_row = 11 + max(len(profile.bank_lines), 3) + 1
    ws.row_dimensions[header_row - 1].height = 12
    headers = (
        (2, "NAZIV", "left"),
        (3, "KOLIČINA", "center"),
        (4, profile.price_header, "right"),
        (5, profile.total_header, "right"),
    )
    header_align = {
        "left": Alignment(horizontal="left", vertical="center"),
        "center": Alignment(horizontal="center", vertical="center"),
        "right": Alignment(horizontal="right", vertical="center"),
    }
    for col, text, align_name in headers:
        _apply(
            ws.cell(header_row, col, text),
            font=_font("Calibri", 9, bold=True, color=WHITE),
            align=header_align[align_name],
            fill=_fill(NAVY),
        )
    ws.row_dimensions[header_row].height = 22

    money_fmt = '#,##0.00 "€"' if profile.symbol else "#,##0.00"
    first_item = header_row + 1
    for index, line in enumerate(invoice.lines):
        row = first_item + index
        zebra = _fill(ZEBRA) if index % 2 else _fill(WHITE)
        border = _border(bottom=True)
        _apply(
            ws.cell(row, 2, line.name),
            font=body_font,
            align=Alignment(horizontal="left", vertical="center", wrap_text=True),
            fill=zebra,
            border=border,
        )
        _apply(
            ws.cell(row, 3, float(line.qty)),
            font=body_font,
            align=Alignment(horizontal="center", vertical="center"),
            fill=zebra,
            border=border,
            fmt="#,##0.##",
        )
        _apply(
            ws.cell(row, 4, float(line.price)),
            font=body_font,
            align=right,
            fill=zebra,
            border=border,
            fmt=money_fmt,
        )
        _apply(
            ws.cell(row, 5, float(line.total)),
            font=bold_font,
            align=right,
            fill=zebra,
            border=border,
            fmt=money_fmt,
        )
        ws.row_dimensions[row].height = 22 if len(line.name) < 55 else 36

    last_item = first_item + len(invoice.lines) - 1
    total_row = last_item + 2
    pay_row = total_row + 1
    _apply(
        ws.cell(total_row, 4, profile.grand_label),
        font=muted_font,
        align=right,
    )
    _apply(
        ws.cell(total_row, 5, float(invoice.total)),
        font=body_font,
        align=right,
        fmt=money_fmt,
    )
    _apply(
        ws.cell(pay_row, 4, "Ukupno za platiti"),
        font=_font("Calibri", 11, bold=True, color=NAVY),
        align=right,
        fill=_fill(TOTAL),
    )
    _apply(
        ws.cell(pay_row, 5, float(invoice.total)),
        font=_font("Calibri", 12, bold=True, color=NAVY),
        align=right,
        fill=_fill(TOTAL),
        fmt=money_fmt,
    )
    ws.row_dimensions[pay_row].height = 24

    cursor = pay_row + 2
    if profile.use_words:
        ws.merge_cells(start_row=cursor, start_column=2, end_row=cursor, end_column=5)
        _apply(
            ws.cell(cursor, 2, amount_in_words(invoice.total, profile.currency)),
            font=_font("Calibri", 10, italic=True, color=INK),
            align=left,
        )
        ws.row_dimensions[cursor].height = 20
        cursor += 2

    for text, is_bold in footer_lines():
        ws.merge_cells(start_row=cursor, start_column=2, end_row=cursor, end_column=5)
        _apply(ws.cell(cursor, 2, text), font=bold_font if is_bold else muted_font, align=left)
        ws.row_dimensions[cursor].height = 18
        cursor += 1

    ws.print_area = f"A1:{get_column_letter(5)}{cursor}"
    ws.page_setup.horizontalCentered = True
    wb.properties.title = f"Račun br. {invoice.number}"
    wb.properties.creator = "Veritas Software"
    wb.properties.subject = f"{profile.title} · {invoice.buyer.name}"
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


_FONTS_READY = False


def _register_fonts() -> tuple[str, str, str]:
    global _FONTS_READY
    regular, bold, italic = "Calibri", "Calibri-Bold", "Calibri-Italic"
    if _FONTS_READY:
        return regular, bold, italic
    fonts = {
        regular: Path(r"C:\Windows\Fonts\calibri.ttf"),
        bold: Path(r"C:\Windows\Fonts\calibrib.ttf"),
        italic: Path(r"C:\Windows\Fonts\calibrii.ttf"),
    }
    if not all(path.exists() for path in fonts.values()):
        fonts = {
            "Arial": Path(r"C:\Windows\Fonts\arial.ttf"),
            "Arial-Bold": Path(r"C:\Windows\Fonts\arialbd.ttf"),
            "Arial-Italic": Path(r"C:\Windows\Fonts\ariali.ttf"),
        }
        regular, bold, italic = "Arial", "Arial-Bold", "Arial-Italic"
    for name, path in fonts.items():
        pdfmetrics.registerFont(TTFont(name, str(path)))
    _FONTS_READY = True
    return regular, bold, italic


def _wrap(text: str, font: str, size: float, max_width: float) -> list[str]:
    words = text.split()
    if not words:
        return [""]
    lines: list[str] = []
    current = words[0]
    for word in words[1:]:
        trial = f"{current} {word}"
        if pdfmetrics.stringWidth(trial, font, size) <= max_width:
            current = trial
        else:
            lines.append(current)
            current = word
    lines.append(current)
    return lines


def write_pdf(invoice: Invoice, path: Path) -> None:
    regular, bold, italic = _register_fonts()
    profile = invoice.profile
    path.parent.mkdir(parents=True, exist_ok=True)
    pdf = canvas.Canvas(str(path), pagesize=A4)
    pdf.setTitle(f"Račun br. {invoice.number}")
    pdf.setAuthor("Veritas Software")

    navy = HexColor("#0E2A47")
    ink = HexColor("#1C2834")
    muted = HexColor("#5C6B7A")
    line = HexColor("#D9E0E8")
    zebra = HexColor("#F5F7FB")
    total_bg = HexColor("#E6E8FF")
    soft = HexColor("#F7F8FB")

    content_w = PAGE_W - 2 * MARGIN
    bottom_limit = 56

    def chrome(continued: bool = False) -> float:
        pdf.setFillColor(navy)
        pdf.rect(0, PAGE_H - 8, PAGE_W, 8, fill=1, stroke=0)
        pdf.setStrokeColor(line)
        pdf.setLineWidth(0.4)
        pdf.line(MARGIN, 40, PAGE_W - MARGIN, 40)
        pdf.setFillColor(muted)
        pdf.setFont(regular, 10)
        pdf.drawCentredString(
            PAGE_W / 2,
            26,
            f"Veritas Software  ·  {SELLER.address}",
        )
        if not continued:
            return PAGE_H - 28
        pdf.setFillColor(navy)
        pdf.setFont(bold, 11)
        pdf.drawString(MARGIN, PAGE_H - 32, f"Račun br. {invoice.number}")
        pdf.setFont(regular, 11)
        pdf.setFillColor(muted)
        pdf.drawRightString(PAGE_W - MARGIN, PAGE_H - 32, "Nastavak")
        return PAGE_H - 52

    def ensure(y: float, needed: float, draw_table_header) -> float:
        if y - needed >= bottom_limit:
            return y
        pdf.showPage()
        y = chrome(continued=True)
        return draw_table_header(y)

    y = chrome(False)
    logo_w, logo_h = 72, 13
    if LOGO_PATH.exists():
        pdf.drawImage(
            ImageReader(str(LOGO_PATH)),
            MARGIN,
            y - logo_h - 4,
            width=logo_w,
            height=logo_h,
            mask="auto",
            preserveAspectRatio=True,
            anchor="sw",
        )
    label = "Račun br."
    number = str(invoice.number)
    pdf.setFillColor(navy)
    pdf.setFont(bold, 28)
    number_w = pdfmetrics.stringWidth(number, bold, 28)
    pdf.drawRightString(PAGE_W - MARGIN, y - 22, number)
    pdf.setFont(bold, 16)
    pdf.drawRightString(PAGE_W - MARGIN - number_w - 10, y - 18, label)

    y -= 40
    pdf.setStrokeColor(line)
    pdf.setLineWidth(1)
    pdf.line(MARGIN, y, PAGE_W - MARGIN, y)
    y -= 22

    gap = 22
    col_w = (content_w - gap) / 2
    right_x = MARGIN + col_w + gap

    def block_height(lines: list[tuple[str, str, float]], width: float) -> float:
        height = 0
        for text, font, size in lines:
            height += len(_wrap(text, font, size, width)) * (size + 3)
        return height + 8

    seller_lines = [
        (SELLER.name, bold, 11),
        (SELLER.address, regular, 10),
        (f"Matični broj: {SELLER.registration}", regular, 10),
        (f"Poreski broj: {SELLER.tax}", regular, 10),
    ]
    buyer_lines = [
        (invoice.buyer.name, bold, 11),
        (invoice.buyer.address, regular, 10),
        (f"Matični broj: {invoice.buyer.registration}", regular, 10),
        (f"Poreski broj: {invoice.buyer.tax}", regular, 10),
    ]
    party_h = max(block_height(seller_lines, col_w), block_height(buyer_lines, col_w)) + 18
    pdf.setFillColor(soft)
    pdf.roundRect(MARGIN, y - party_h, col_w, party_h, 6, fill=1, stroke=0)
    pdf.roundRect(right_x, y - party_h, col_w, party_h, 6, fill=1, stroke=0)

    def paint_party(x: float, title: str, lines: list[tuple[str, str, float]]) -> None:
        cursor = y - 16
        pdf.setFillColor(navy)
        pdf.setFont(bold, 10)
        pdf.drawString(x + 10, cursor, title)
        cursor -= 16
        for text, font, size in lines:
            for wrapped in _wrap(text, font, size, col_w - 20):
                pdf.setFont(font, size)
                pdf.setFillColor(navy if font == bold else ink)
                pdf.drawString(x + 10, cursor, wrapped)
                cursor -= size + 3

    paint_party(MARGIN, "IZDAVALAC", seller_lines)
    paint_party(right_x, "KUPAC", buyer_lines)
    y -= party_h + 18

    bank_wrapped: list[str] = []
    for entry in profile.bank_lines:
        bank_wrapped.extend(_wrap(entry, regular, 10, col_w))
    meta = meta_lines(invoice)
    meta_wrapped: list[tuple[str, bool]] = []
    for text, is_bold in meta:
        font = bold if is_bold else regular
        for wrapped in _wrap(text, font, 10, col_w):
            meta_wrapped.append((wrapped, is_bold))
    info_h = max(len(bank_wrapped), len(meta_wrapped)) * 13
    cursor = y
    for wrapped in bank_wrapped:
        pdf.setFont(regular, 10)
        pdf.setFillColor(ink)
        pdf.drawString(MARGIN, cursor, wrapped)
        cursor -= 13
    cursor = y
    for wrapped, is_bold in meta_wrapped:
        pdf.setFont(bold if is_bold else regular, 10)
        pdf.setFillColor(navy if is_bold else ink)
        pdf.drawString(right_x, cursor, wrapped)
        cursor -= 13
    y -= info_h + 18

    qty_w, price_w, total_w = 86, 96, 110
    name_w = content_w - qty_w - price_w - total_w
    cols = (
        (MARGIN, name_w, "NAZIV", "left"),
        (MARGIN + name_w, qty_w, "KOLIČINA", "center"),
        (MARGIN + name_w + qty_w, price_w, profile.price_header, "right"),
        (MARGIN + name_w + qty_w + price_w, total_w, profile.total_header, "right"),
    )

    def table_header(top: float) -> float:
        pdf.setFillColor(navy)
        pdf.roundRect(MARGIN, top - 28, content_w, 28, 3, fill=1, stroke=0)
        pdf.rect(MARGIN, top - 28, content_w, 10, fill=1, stroke=0)
        pdf.setFillColor(white)
        pdf.setFont(bold, 11)
        for x, width, text, align in cols:
            baseline = top - 18
            if align == "left":
                pdf.drawString(x + 8, baseline, text)
            elif align == "center":
                pdf.drawCentredString(x + width / 2, baseline, text)
            else:
                pdf.drawRightString(x + width - 8, baseline, text)
        return top - 28

    y = table_header(y)

    def row_lines(text: str) -> list[str]:
        return _wrap(text, regular, 12, name_w - 16)

    for index, item in enumerate(invoice.lines):
        wrapped = row_lines(item.name)
        row_h = max(32, 14 + len(wrapped) * 16)
        y = ensure(y, row_h + 4, table_header)
        if index % 2 == 1:
            pdf.setFillColor(zebra)
            pdf.rect(MARGIN, y - row_h, content_w, row_h, fill=1, stroke=0)
        pdf.setStrokeColor(line)
        pdf.setLineWidth(0.3)
        pdf.line(MARGIN, y - row_h, MARGIN + content_w, y - row_h)
        text_y = y - 20
        pdf.setFillColor(ink)
        pdf.setFont(regular, 12)
        for wrapped_line in wrapped:
            pdf.drawString(MARGIN + 8, text_y, wrapped_line)
            text_y -= 16
        mid = y - row_h / 2 - 4
        pdf.setFont(regular, 12)
        pdf.drawCentredString(MARGIN + name_w + qty_w / 2, mid, fmt_qty(item.qty))
        pdf.drawRightString(MARGIN + name_w + qty_w + price_w - 8, mid, fmt_money(item.price))
        pdf.setFont(bold, 12)
        pdf.drawRightString(PAGE_W - MARGIN - 8, mid, fmt_money(item.total))
        y -= row_h

    y = ensure(y, 78, table_header)
    y -= 16
    box_w = 250
    box_x = PAGE_W - MARGIN - box_w
    amount = fmt_money(invoice.total)
    pay_amount = f"{amount} {profile.symbol}".strip()
    pdf.setFillColor(muted)
    pdf.setFont(regular, 11)
    pdf.drawString(box_x, y, profile.grand_label)
    pdf.setFillColor(ink)
    pdf.setFont(regular, 12)
    pdf.drawRightString(PAGE_W - MARGIN, y, amount)
    y -= 28
    pdf.setFillColor(total_bg)
    pdf.roundRect(box_x, y - 6, box_w, 26, 4, fill=1, stroke=0)
    pdf.setFillColor(navy)
    pdf.setFont(bold, 12)
    pdf.drawString(box_x + 10, y + 2, "Ukupno za platiti")
    pdf.setFont(bold, 14)
    pdf.drawRightString(PAGE_W - MARGIN - 10, y + 1, pay_amount)
    y -= 28

    if profile.use_words:
        words = amount_in_words(invoice.total, profile.currency)
        for wrapped in _wrap(words, italic, 10, content_w):
            y = ensure(y, 16, table_header)
            pdf.setFillColor(ink)
            pdf.setFont(italic, 10)
            pdf.drawString(MARGIN, y, wrapped)
            y -= 14
        y -= 8

    y -= 6
    for text, is_bold in footer_lines():
        y = ensure(y, 16, table_header)
        pdf.setFillColor(navy if is_bold else muted)
        pdf.setFont(bold if is_bold else regular, 11)
        pdf.drawString(MARGIN, y, text)
        y -= 14

    pdf.save()
